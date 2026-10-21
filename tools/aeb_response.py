"""AEB 制动响应时间统计：批量读 blf / asc，输出 Excel。

起点：ACC_AEBActive 从 0（not Active）跳变为 1（Active）的第一帧。
终点：之后每收到一帧 ESP_LongAccel，就和它之前最近一帧 ACC_AEBTargetDeceleration 比较，
     |ESP_LongAccel| 第一次大于 |目标减速度| 的时刻。
     口径A（字面）：激活后第一次满足就算。
     口径B（推荐）：|目标| 达到本次峰值的一半（--ramp）以后才判断，且要求 |减速度| 先低于目标、再第一次超过目标。
     （目标减速度激活时从 -0.4～-2 逐步爬升到 -8/-9，激活前车辆已有少量减速度时，口径A 会得到几 ms 的无效结果。）
一个日志里有多次激活时，每次单独一行。

依赖：pip install python-can openpyxl cryptography
用法：
  export VEHDIAG_SRC_KEY='原始文件密码'
  python tools/aeb_response.py 报文目录或文件... --model C363H [--out 结果.xlsx]
  可选：--active GW_244.ACC_AEBActive --target GW_244.ACC_AEBTargetDeceleration
        --accel GW_15D.ESP_LongAccel --speed IBCU_17A.ESP_VehicleSpeed   （“报文名.信号名”，信号名唯一时可省报文名）
"""
import argparse, bisect, glob, json, os, re, sys, warnings

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import analyze as A  # noqa: E402

warnings.filterwarnings('ignore')
LOG_EXT = ('.blf', '.asc')


def pick(db, spec):
    """'报文名.信号名' 或 '信号名' → (msg, sig)。"""
    mn, _, sn = spec.rpartition('.')
    hits = [(m, s) for m in db['msgs'] for s in m['sigs'] if s['n'] == sn and (not mn or m['name'] == mn)]
    if not hits:
        sys.exit(f'DBC 里没有信号 {spec}')
    if len(hits) > 1:
        sys.exit(f'信号 {sn} 在多个报文里：' + '、'.join(m['name'] for m, _ in hits) + '，请写成 报文名.信号名')
    return hits[0]


def collect(paths):
    out = []
    for p in paths:
        if os.path.isdir(p):
            out += [f for f in sorted(glob.glob(os.path.join(p, '**', '*'), recursive=True)) if f.lower().endswith(LOG_EXT)]
        else:
            out.append(p)
    return out


def analyze_file(path, db, a):
    fr = A.read_log(path)
    idx = A.by_id(fr)
    act = A.decode(fr, idx, *pick(db, a.active))
    tgt = A.decode(fr, idx, *pick(db, a.target))
    acc = A.decode(fr, idx, *pick(db, a.accel))
    spd = A.decode(fr, idx, *pick(db, a.speed)) if a.speed else []
    if not act or not tgt or not acc:
        return [], '缺少信号帧（' + '、'.join(n for n, v in (('Active', act), ('目标减速度', tgt), ('纵向加速度', acc)) if not v) + '）'
    tt, st = [t for t, _, _ in tgt], [t for t, _, _ in spd]

    def latest(series, times, t):
        k = bisect.bisect_right(times, t) - 1
        return series[k][2] if k >= 0 else None

    events = []
    for i in range(1, len(act)):
        if act[i - 1][1] == 0 and act[i][1] == 1:
            t0, j = act[i][0], i
            while j < len(act) and act[j][1] == 1:
                j += 1
            tend = act[j][0] if j < len(act) else act[-1][0]
            peak = max((abs(v) for t, _, v in tgt if t0 <= t <= tend), default=0)
            hit_a = hit_b = None
            armed = False
            for t, _, v in acc:
                if t < t0 or t > tend:
                    continue
                lt = latest(tgt, tt, t)
                if lt is None:
                    continue
                if hit_a is None and abs(v) > abs(lt):
                    hit_a = (t, v, lt)
                if abs(lt) < a.ramp * peak:                 # 目标还在爬升初期，不判
                    continue
                if abs(v) <= abs(lt):
                    armed = True
                if hit_b is None and armed and abs(v) > abs(lt):
                    hit_b = (t, v, lt)
                if hit_a and hit_b:
                    break
            a0 = latest(acc, [t for t, _, _ in acc], t0)
            detail = [(t, v, latest(tgt, tt, t), next((r for tq, r, _ in reversed(act) if tq <= t), 0))
                      for t, _, v in acc if t0 - 0.1 <= t <= t0 + a.window]
            events.append(dict(t0=t0, tend=tend, A=hit_a, B=hit_b, detail=detail, a0=a0,
                               prev_gap=t0 - events[-1]['tend'] if events else None,
                               v0=latest(spd, st, t0) if spd else None))
    return events, '' if events else '没有 ACC_AEBActive 0→1 跳变'


def kph_label(name):
    m = re.search(r'(\d+)\s*KPH', name, re.I)
    return int(m.group(1)) if m else None


def write_xlsx(rows, out, a):
    """所有结果直接写数值（不用公式）：手机和预览软件不会计算公式，写公式会显示为空。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.chart import ScatterChart, Reference, Series
    F = 'Arial'
    hf, hfill = Font(name=F, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E79')
    body, bold = Font(name=F), Font(name=F, bold=True)
    key = PatternFill('solid', fgColor='DDEBF7')
    thin = Side(style='thin', color='BFBFBF')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def head(ws, r, names, h=48):
        for c, n in enumerate(names, 1):
            x = ws.cell(r, c, n)
            x.font, x.fill, x.border = hf, hfill, box
            x.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        ws.row_dimensions[r].height = h

    wb = Workbook()
    ws = wb.active
    ws.title = '汇总'
    ws['A1'] = 'AEB 制动响应时间统计（ACC_AEBActive 激活 → |ESP_LongAccel| 超过目标减速度）'
    ws['A1'].font = Font(name=F, bold=True, size=13)
    ws['A2'] = f'DBC：{a.dbc_name}；时间为相对日志第一帧的秒数；时间差 = (终点 - 起点) × 1000。蓝底列为主要结果。'
    ws['A2'].font = Font(name=F, color='595959')
    r0 = 4
    head(ws, r0, ['日志文件', '序号', '文件名车速 (km/h)', '起点实际车速 (km/h)', '起点时刻 (s)',
                  '终点时刻 (s)\n【口径B·推荐】', '时间差 (ms)\n【口径B·推荐】', '终点 ESP_LongAccel (m/s²)', '终点时目标减速度 (m/s²)',
                  '终点时刻 (s)\n【口径A·字面】', '时间差 (ms)\n【口径A·字面】', '终点 ESP_LongAccel (m/s²)', '终点时目标减速度 (m/s²)',
                  'AEB 激活持续 (s)', '备注'])
    r = r0
    for name, k, e, msg in rows:
        r += 1
        if e is None:
            vals = [name, None, kph_label(name)] + [None] * 11 + [msg]
        else:
            ha, hb = e['A'], e['B']
            notes = []
            if e['a0'] is not None and abs(e['a0']) >= a.braking:
                gap = f"（距上次 AEB 结束 {e['prev_gap']:.1f} s）" if e['prev_gap'] is not None else ''
                notes.append(f"激活时车辆已在制动（|减速度| {abs(e['a0']):.1f}）{gap}，结果不代表从零开始的制动响应")
            if hb is None:
                notes.append('激活期间 |减速度| 未超过目标')
            elif ha and ha[0] != hb[0]:
                notes.append(f'激活初期目标值很小（{ha[2]:.1f}），|减速度| 本身的波动就超过了它，口径A 无参考意义')
            note = '；'.join(notes)
            vals = [name, k, kph_label(name), None if e['v0'] is None else round(e['v0'], 1), round(e['t0'], 6),
                    hb and round(hb[0], 6), hb and round((hb[0] - e['t0']) * 1000, 1), hb and round(hb[1], 3), hb and round(hb[2], 3),
                    ha and round(ha[0], 6), ha and round((ha[0] - e['t0']) * 1000, 1), ha and round(ha[1], 3), ha and round(ha[2], 3),
                    round(e['tend'] - e['t0'], 3), note]
        for c, v in enumerate(vals, 1):
            x = ws.cell(r, c, v)
            x.font, x.border = body, box
            x.alignment = Alignment(horizontal='left' if c in (1, 15) else 'center', vertical='center', wrap_text=c == 15)
            if c in (6, 7):
                x.fill = key
        ws.cell(r, 7).font = bold
        for c in (5, 6, 10):
            ws.cell(r, c).number_format = '0.0000'
        for c in (7, 11):
            ws.cell(r, c).number_format = '0.0'
        for c in (8, 9, 12, 13, 14):
            ws.cell(r, c).number_format = '0.000'
    last = r
    for i, (lab, fn) in enumerate((('平均值', 'AVERAGE'), ('最小值', 'MIN'), ('最大值', 'MAX'))):
        rr = last + 1 + i
        ws.cell(rr, 1, lab).font = bold
        for col in ('G', 'K'):
            x = ws[f'{col}{rr}']
            vs = [ws[f'{col}{k}'].value for k in range(r0 + 1, last + 1)]
            vs = [v for v in vs if isinstance(v, (int, float))]
            x.value = round({'AVERAGE': sum(vs) / len(vs) if vs else 0, 'MIN': min(vs, default=0), 'MAX': max(vs, default=0)}[fn], 1) if vs else ''
            x.font, x.number_format = bold, '0.0'
            x.alignment = Alignment(horizontal='center')
            if col == 'G':
                x.fill = key
        for c in range(1, 16):
            ws.cell(rr, c).border = box
    notes = [
        '口径说明',
        '起点：ACC_AEBActive 从 0 "not Active" 跳变为 1 "Active" 的第一帧时刻；一个日志里多次激活时每次一行（序号）。',
        '终点：每收到一帧 ESP_LongAccel，就和它之前最近一帧 ACC_AEBTargetDeceleration 比较；目标值是负数，两边都取绝对值。',
        '口径A（字面）：激活后第一次满足 |ESP_LongAccel| > |ACC_AEBTargetDeceleration|。',
        f'口径B（推荐）：只在 |目标减速度| 达到本次激活峰值的 {a.ramp:.0%} 以上之后判断，并要求 |减速度| 先低于目标、再第一次超过目标。'
        '目标减速度在激活时从 -0.4～-2 逐步爬升，这段时间里 |减速度| 本身的波动就可能超过目标，口径A 会得到几 ms 的无效结果。两种口径相同时说明没有这种情况。',
        f'备注里“激活时车辆已在制动”指激活时刻 |ESP_LongAccel| ≥ {a.braking:g} m/s²（多为 AEB 刚结束又再次激活），这类结果不代表从零开始的制动响应。',
        '分辨率：结果取帧时刻、不插值，精度约等于 ESP_LongAccel 的报文周期（约 10 ms）。',
        '“起点实际车速”取 ESP_VehicleSpeed 在起点时刻的值，供核对试验车速。各分页为每次激活前后的曲线数据，黄色行为口径B 终点帧。',
    ]
    n = last + 5
    for i, t in enumerate(notes):
        x = ws.cell(n + i, 1, t)
        x.font = bold if i == 0 else body
        x.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=n + i, start_column=1, end_row=n + i, end_column=15)
        ws.row_dimensions[n + i].height = 18 if len(t) < 90 else 36
    for i, w in enumerate([24, 6, 10, 11, 13, 14, 13, 13, 13, 14, 13, 13, 13, 11, 46]):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'B5'

    used = set()
    for name, k, e, _ in rows:
        if e is None:
            continue
        base = (f'{kph_label(name)}kph' if kph_label(name) is not None else os.path.splitext(name)[0])[:25]
        title = base + (f'_{k}' if k > 1 else '')
        while title in used:
            title += '_'
        used.add(title)
        d = wb.create_sheet(title)
        hb = e['B']
        d['A1'] = f"{name}  起点 t0 = {e['t0']:.4f} s；口径B 终点 " + (f"+{(hb[0] - e['t0']) * 1000:.1f} ms" if hb else '未达到')
        d['A1'].font = Font(name=F, bold=True, size=12)
        head(d, 3, ['相对起点 (ms)', '日志时刻 (s)', 'ESP_LongAccel (m/s²)', 'ACC_AEBTargetDeceleration\n最近一帧 (m/s²)',
                    '|ESP_LongAccel|', '|目标减速度|', 'ACC_AEBActive', '是否超过目标'], 34)
        for i, (t, v, tg, ac) in enumerate(e['detail']):
            rr = 4 + i
            row = [round((t - e['t0']) * 1000, 1), round(t, 6), round(v, 4), None if tg is None else round(tg, 4),
                   round(abs(v), 4), None if tg is None else round(abs(tg), 4), ac,
                   '是' if ac == 1 and tg is not None and abs(v) > abs(tg) else '']
            for c, val in enumerate(row, 1):
                x = d.cell(rr, c, val)
                x.font, x.border = body, box
                x.alignment = Alignment(horizontal='center')
            d.cell(rr, 1).number_format = '0.0'
            d.cell(rr, 2).number_format = '0.0000'
            for c in (3, 4, 5, 6):
                d.cell(rr, c).number_format = '0.000'
            if hb and abs(t - hb[0]) < 1e-9:
                for c in range(1, 9):
                    d.cell(rr, c).fill = PatternFill('solid', fgColor='FFE699')
        end = 3 + len(e['detail'])
        for c, w in enumerate([13, 13, 16, 22, 15, 13, 14, 12], 1):
            d.column_dimensions[chr(64 + c)].width = w
        d.freeze_panes = 'A4'
        ch = ScatterChart()
        ch.title, ch.style, ch.height, ch.width = '|减速度| 与 |目标减速度|', 2, 9, 18
        ch.x_axis.title, ch.y_axis.title = '相对起点 (ms)', 'm/s²'
        ch.x_axis.delete = ch.y_axis.delete = False
        xs = Reference(d, min_col=1, min_row=4, max_row=end)
        for col, color in ((5, '1F4E79'), (6, 'C00000')):
            s = Series(Reference(d, min_col=col, min_row=3, max_row=end), xs, title_from_data=True)
            s.marker.symbol = 'none'
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 19000
            ch.series.append(s)
        d.add_chart(ch, 'J3')
    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description='AEB 制动响应时间统计')
    ap.add_argument('logs', nargs='+', help='blf / asc 文件或目录（目录会递归查找）')
    ap.add_argument('--src-key', default=os.environ.get('VEHDIAG_SRC_KEY'))
    ap.add_argument('--model', help='车型（项目代号如 C363H，或市场名），决定 DBC')
    ap.add_argument('--dbc', help='直接指定 DBC 名')
    ap.add_argument('--active', default='GW_244.ACC_AEBActive')
    ap.add_argument('--target', default='GW_244.ACC_AEBTargetDeceleration')
    ap.add_argument('--accel', default='GW_15D.ESP_LongAccel')
    ap.add_argument('--speed', default='IBCU_17A.ESP_VehicleSpeed', help='起点车速信号，填空字符串则不取')
    ap.add_argument('--window', type=float, default=0.8, help='分页曲线数据取起点后多少秒（默认 0.8）')
    ap.add_argument('--ramp', type=float, default=0.5, help='口径B：|目标| 达到本次峰值的多少比例后才开始判断（默认 0.5）')
    ap.add_argument('--braking', type=float, default=3.0, help='激活时 |减速度| 超过多少 m/s² 视为已在制动，备注提示（默认 3）')
    ap.add_argument('--out', default='AEB制动响应时间.xlsx')
    a = ap.parse_args()
    if not a.src_key:
        ap.error('需要原始文件密码 --src-key（或环境变量 VEHDIAG_SRC_KEY）')
    dbcs, _ = A.load_data(a.src_key)
    dbc_map = json.load(open(os.path.join(A.ROOT, 'tools', 'dbc_map.json'), encoding='utf-8'))
    name = a.dbc or dbc_map.get(A.build.market_name(a.model or ''))
    if name not in dbcs:
        sys.exit('请用 --model 指定车型（可选：' + '、'.join(A.build.code_name(k) for k in dbc_map if not k.startswith('_')) + '）或 --dbc 指定 DBC')
    db, a.dbc_name = dbcs[name], name
    files = collect(a.logs)
    if not files:
        sys.exit('没有找到 blf / asc 文件')
    rows = []
    for f in files:
        evs, msg = analyze_file(f, db, a)
        base = os.path.basename(f)
        if not evs:
            rows.append((base, None, None, msg))
            print(f'{base}: {msg}')
        for k, e in enumerate(evs, 1):
            rows.append((base, k, e, ''))
            fmt = lambda h: f"{(h[0] - e['t0']) * 1000:.1f} ms" if h else '未达到'
            print(f"{base} #{k}: 起点 {e['t0']:.4f} s  口径B {fmt(e['B'])}  口径A {fmt(e['A'])}")
    rows.sort(key=lambda x: (kph_label(x[0]) is None, kph_label(x[0]) or 0, x[0], x[1] or 0))
    write_xlsx(rows, a.out, a)
    print(f'已生成 {a.out}')


if __name__ == '__main__':
    main()
